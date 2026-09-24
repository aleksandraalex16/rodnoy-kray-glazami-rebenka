from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("Проект_конкурса_Родной_край_глазами_ребенка.md")
OUTPUT = Path("Проект_конкурса_Родной_край_глазами_ребенка.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_inline_markdown(paragraph, text):
    pieces = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(piece)


def clean_table_cell(value):
    return value.strip().replace("**", "").replace("`", "")


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(1.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = None

styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].paragraph_format.space_before = Pt(8)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
i = 0
title_count = 0
while i < len(lines):
    line = lines[i].rstrip()
    stripped = line.strip()

    if not stripped:
        i += 1
        continue

    if stripped == "---":
        next_content = next((item.strip() for item in lines[i + 1:] if item.strip()), "")
        if next_content.startswith("## ") or next_content.startswith("# Приложение"):
            doc.add_section(WD_SECTION.NEW_PAGE)
        i += 1
        continue

    if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            if not re.match(r"^\|[\s:|-]+\|$", lines[i].strip()):
                rows.append([clean_table_cell(x) for x in lines[i].strip().strip("|").split("|")])
            i += 1
        if rows:
            columns = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=columns)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for row_index, values in enumerate(rows):
                for column_index in range(columns):
                    cell = table.cell(row_index, column_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell)
                    value = values[column_index] if column_index < len(values) else ""
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1
                        for run in paragraph.runs:
                            run.font.name = "Times New Roman"
                            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                            run.font.size = Pt(9)
                            if row_index == 0:
                                run.bold = True
                    if row_index == 0:
                        set_cell_shading(cell, "D9EAD3")
        continue

    if stripped.startswith("# "):
        title_count += 1
        text = stripped[2:]
        if text.startswith("Приложение"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(text)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        else:
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(text)
        i += 1
        continue

    if stripped.startswith("## "):
        text = stripped[3:]
        if title_count and not re.match(r"\d+\.", text):
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(text)
        else:
            doc.add_paragraph(text, style="Heading 1")
        i += 1
        continue

    if stripped.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        add_inline_markdown(paragraph, stripped[2:])
        i += 1
        continue

    paragraph = doc.add_paragraph()
    if re.match(r"^\d+\.\d+\.", stripped):
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    if stripped.startswith("**Организатор:") or stripped.startswith("**Сроки проведения:") or stripped.startswith("**Место проведения:") or stripped.startswith("**Формат:"):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_markdown(paragraph, stripped.replace("  ", ""))
    i += 1

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))

doc.core_properties.title = "Республиканский конкурс детского творчества «Родной край — глазами ребенка»"
doc.core_properties.subject = "Проект и положение о конкурсе"
doc.core_properties.author = "Организатор конкурса"
doc.save(OUTPUT)
print(OUTPUT)
